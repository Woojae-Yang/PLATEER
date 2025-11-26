import openpyxl
import os

class LoginData:

    @staticmethod
    def get_excel_data(data_num):

        # LoginData.py가 있는 TestData 폴더 절대경로
        current_dir = os.path.dirname(os.path.abspath(__file__))

        # 엑셀 파일 경로 (같은 폴더에 위치)
        excel_path = os.path.join(current_dir, "GroobeeLoginData.xlsx")

        # 엑셀 파일 로드
        workbook = openpyxl.load_workbook(excel_path)
        sheet = workbook.active

        result = {}

        # 엑셀 데이터 읽기
        for i in range(2, sheet.max_row + 1):
            if str(sheet.cell(row=i, column=1).value) == str(data_num):
                for j in range(2, sheet.max_column + 1):
                    key = sheet.cell(row=1, column=j).value
                    value = sheet.cell(row=i, column=j).value
                    result[key] = value
                break

        workbook.close()
        return [result]